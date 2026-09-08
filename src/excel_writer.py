"""
excel_writer.py  (Tarea 1)
--------------------------
Registra en un Excel local los CONTENEDORES cerrados en las últimas 24h junto
con sus Task hijas. Un contenedor puede ser una HDU, un Enabler, un Objetivo
IT, etc.: todo lo que sea padre de una Task. Idempotente: la clave única
"{contenedor_id}-{task_id}" evita duplicar filas al re-ejecutar.
"""
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNAS = [
    "clave", "cont_id", "cont_tipo", "cont_titulo", "cont_estado", "cont_cerrado_el",
    "tarea_id", "tarea_titulo", "tarea_estado", "tarea_asignado",
    "tarea_horas", "tarea_horas_reales", "tarea_fecha_inicio", "tarea_fecha_fin",
]


def construir_registros(contenedores_cerrados, tareas_por_padre):
    filas = []
    for c in contenedores_cerrados:
        hijas = tareas_por_padre.get(c["id"], [])
        ref = c["closed_date"] or c["changed_date"]
        cerrado_el = ref.strftime("%Y-%m-%d %H:%M") if ref else ""
        base = {
            "cont_id": c["id"], "cont_tipo": c["tipo"], "cont_titulo": c["titulo"],
            "cont_estado": c["estado"], "cont_cerrado_el": cerrado_el,
        }
        if not hijas:
            filas.append({**base, "clave": f"{c['id']}-NA",
                          "tarea_id": "", "tarea_titulo": "(sin Task hijas)",
                          "tarea_estado": "", "tarea_asignado": "", "tarea_horas": "",
                          "tarea_horas_reales": "", "tarea_fecha_inicio": "",
                          "tarea_fecha_fin": ""})
            continue
        for t in hijas:
            filas.append({
                **base, "clave": f"{c['id']}-{t['id']}",
                "tarea_id": t["id"], "tarea_titulo": t["titulo"],
                "tarea_estado": t["estado"], "tarea_asignado": t["asignado"],
                "tarea_horas": t["horas"], "tarea_horas_reales": t["horas_reales"],
                "tarea_fecha_inicio": t["fecha_inicio"].isoformat() if t["fecha_inicio"] else "",
                "tarea_fecha_fin": t["fecha_fin"].isoformat() if t["fecha_fin"] else "",
            })
    return filas


def upsert_excel(ruta, registros):
    """Combina los registros nuevos con lo ya existente (upsert por 'clave')."""
    nuevos = pd.DataFrame(registros, columns=COLUMNAS)

    if os.path.exists(ruta):
        try:
            existentes = pd.read_excel(ruta, dtype=str)
        except Exception:
            existentes = pd.DataFrame(columns=COLUMNAS)
        df = pd.concat([existentes, nuevos.astype(str)], ignore_index=True)
    else:
        df = nuevos.astype(str)

    df = df.drop_duplicates(subset=["clave"], keep="last").reset_index(drop=True)
    df = df.sort_values(["cont_cerrado_el", "cont_id"], ascending=[False, True])

    _escribir_con_formato(df, ruta)
    return len(nuevos), len(df)


def _escribir_con_formato(df, ruta):
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Contenedores Cerrados")
        ws = writer.sheets["Contenedores Cerrados"]
        encabezado = Font(bold=True, color="FFFFFF")
        relleno = PatternFill("solid", fgColor="B87333")  # cobre
        for celda in ws[1]:
            celda.font = encabezado
            celda.fill = relleno
            celda.alignment = Alignment(horizontal="center")
        anchos = {"A": 14, "B": 9, "C": 14, "D": 40, "E": 12, "F": 18, "G": 9,
                  "H": 40, "I": 14, "J": 24, "K": 8, "L": 10, "M": 14, "N": 14}
        for col, w in anchos.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
